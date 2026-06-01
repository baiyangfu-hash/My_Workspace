# -*- coding: utf-8 -*-
"""Excel导出器 - FB/FC接口变量表导出为Excel格式"""
import os
import re
from pathlib import Path
from typing import List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

from src.exporters.fb_interface_template import FBInterfaceExport, FBInterfaceVariable
from src.parsers.st_parser import STParser, VarCategory
from src.utils.logger import setup_logger

logger = setup_logger(__name__)

CATEGORY_MAP = {
    VarCategory.VAR_INPUT: "IN",
    VarCategory.VAR_OUTPUT: "OUT",
    VarCategory.VAR_IN_OUT: "IN_OUT",
    VarCategory.VAR: "VAR",
    VarCategory.VAR_STATIC: "VAR",
    VarCategory.VAR_TEMP: "VAR",
}

HEADER_ROW = ["序号", "类别", "名称", "数据类型", "隐藏", "初始值", "掉电保持", "注释"]

HEADER_FONT = Font(bold=True, size=11, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

DATA_ALIGNMENT_CENTER = Alignment(horizontal="center", vertical="center")
DATA_ALIGNMENT_LEFT = Alignment(horizontal="left", vertical="center")

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

EVEN_ROW_FILL = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
ODD_ROW_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

RE_POU_HEADER = re.compile(
    r"(FUNCTION_BLOCK|FUNCTION)\s+(?P<name>\w+)",
    re.IGNORECASE,
)

RE_VAR_BLOCK = re.compile(
    r"(VAR(?:_(?:INPUT|OUTPUT|IN_OUT|STATIC|TEMP))?)"
    r"(?:\s+(?P<retain>RETAIN|NON_RETAIN|PERSISTENT))?"
    r"\s*\n"
    r"(?P<body>.*?)"
    r"END_VAR",
    re.IGNORECASE | re.DOTALL,
)

RE_VAR_LINE = re.compile(
    r"(?P<name>\w+)"
    r"\s*:\s*"
    r"(?P<type>\w+(?:\s*\([^)]*\))?(?:\s*ARRAY\s+\[.*?\])?(?:\s*OF\s+\w+)?)"
    r"(?:\s*:=\s*(?P<init>[^;]*?))?"
    r"\s*;"
    r"\s*(?:\(\*\s*(?P<cmt1>(?:[^*]|\*(?!))*)\*\))?"
    r"\s*(?://\s*(?P<cmt2>.*))?$",
    re.IGNORECASE | re.MULTILINE,
)


class ExcelExporter:
    """FB接口变量Excel导出器"""

    @classmethod
    def _parse_scl_source(cls, scl_source: str, target_fb_name: str = None) -> List[FBInterfaceExport]:
        results = []

        pou_match = RE_POU_HEADER.search(scl_source)
        if not pou_match:
            logger.warning("未找到POU定义")
            return results

        fb_type = pou_match.group(1).upper()
        fb_name = pou_match.group("name").strip()

        if target_fb_name and fb_name != target_fb_name:
            return results

        export = FBInterfaceExport(fb_name=fb_name, fb_type=fb_type)
        seq = 1

        for var_block in RE_VAR_BLOCK.finditer(scl_source):
            cat_raw = var_block.group(1).upper()
            retain_kw = var_block.group("retain")
            body = var_block.group("body")

            try:
                cat_enum = VarCategory(cat_raw)
            except ValueError:
                cat_enum = VarCategory.VAR

            category = CATEGORY_MAP.get(cat_enum, "VAR")

            if retain_kw and retain_kw.upper() == "RETAIN":
                retain_str = "保持"
            elif retain_kw and retain_kw.upper() == "PERSISTENT":
                retain_str = "保持"
            else:
                retain_str = "不保持"

            for line_match in RE_VAR_LINE.finditer(body):
                var_name = line_match.group("name").strip()
                data_type = line_match.group("type") or "UNKNOWN"
                init_val = (line_match.group("init") or "").strip()
                cmt1 = (line_match.group("cmt1") or "").strip()
                cmt2 = (line_match.group("cmt2") or "").strip()
                comment = cmt1 or cmt2

                if data_type.startswith("STRUCT") or data_type.startswith("ST_"):
                    pass
                elif re.match(r"^FB_\w+$", data_type) or re.match(r"^FC_\w+$", data_type):
                    pass

                variable = FBInterfaceVariable(
                    seq_no=seq,
                    category=category,
                    name=var_name,
                    data_type=data_type.strip(),
                    hidden="OFF",
                    initial_value=init_val,
                    retain_type=retain_str,
                    comment=comment,
                )
                export.variables.append(variable)
                seq += 1

        export.recalculate_totals()
        results.append(export)
        logger.info(f"解析完成: {fb_name}, 共 {len(export.variables)} 个变量")
        return results

    @classmethod
    def _write_sheet(cls, ws, export: FBInterfaceExport):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADER_ROW))
        title_cell = ws.cell(row=1, column=1, value=f"{export.fb_name} 接口变量表")
        title_cell.font = Font(bold=True, size=14)
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        for col_idx, header in enumerate(HEADER_ROW, start=1):
            cell = ws.cell(row=2, column=col_idx, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGNMENT
            cell.border = THIN_BORDER

        for row_idx, var in enumerate(export.variables, start=3):
            row_data = var.to_row()
            fill = EVEN_ROW_FILL if (row_idx % 2 == 0) else ODD_ROW_FILL

            for col_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = THIN_BORDER
                cell.fill = fill

                if col_idx in (1, 2, 5, 7):
                    cell.alignment = DATA_ALIGNMENT_CENTER
                else:
                    cell.alignment = DATA_ALIGNMENT_LEFT

        cls._auto_fit_columns(ws)

        ws.freeze_panes = "A3"

    @classmethod
    def _auto_fit_columns(cls, ws):
        col_widths = {1: 6, 2: 8, 3: 24, 4: 14, 5: 8, 6: 20, 7: 10, 8: 40}
        for col_num, width in col_widths.items():
            ws.column_dimensions[get_column_letter(col_num)].width = width

    @classmethod
    def export_fb_interface(cls, scl_source: str, fb_name: str, output_path: str) -> str:
        exports = cls._parse_scl_source(scl_source, target_fb_name=fb_name)
        if not exports:
            raise ValueError(f"未在源码中找到FB/FC: {fb_name}")

        wb = Workbook()
        ws = wb.active
        ws.title = fb_name

        cls._write_sheet(ws, exports[0])

        output_path = Path(output_path).resolve()
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(output_path))

        logger.info(f"FB接口变量表已导出: {output_path}")
        return str(output_path)

    @classmethod
    def export_fb_from_file(cls, scl_file_path: str, output_path: str = None) -> str:
        scl_file = Path(scl_file_path)
        if not scl_file.exists():
            raise FileNotFoundError(f"SCL文件不存在: {scl_file_path}")

        source_text = scl_file.read_text(encoding="utf-8")

        exports = cls._parse_scl_source(source_text)
        if not exports:
            raise ValueError(f"文件中未找到POU定义: {scl_file_path}")

        export = exports[0]
        export.source_file = str(scl_file.resolve())

        if output_path is None:
            output_path = str(scl_file.with_suffix(".xlsx"))

        wb = Workbook()
        ws = wb.active
        ws.title = export.fb_name

        cls._write_sheet(ws, export)

        out = Path(output_path)
        if out.is_dir() or not out.suffix:
            out = out / f"{export.fb_name}.xlsx"
        out.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(out))

        logger.info(f"FB接口变量表已导出: {out}")
        return str(out)

    @classmethod
    def export_all_fbs(
        cls,
        project_path: str,
        output_dir: str,
        include_syslib: bool = False,
    ) -> List[str]:
        project = Path(project_path)
        if not project.exists():
            raise FileNotFoundError(f"项目目录不存在: {project_path}")

        output = Path(output_dir)
        output.mkdir(parents=True, exist_ok=True)

        pattern = "**/*.scl"
        scl_files = sorted(project.glob(pattern), key=lambda p: p.name.lower())

        exported_paths = []

        for scl_file in scl_files:
            if not include_syslib and "syslib" in scl_file.parts:
                continue

            try:
                source_text = scl_file.read_text(encoding="utf-8")
                exports = cls._parse_scl_source(source_text)

                for exp in exports:
                    exp.source_file = str(scl_file.resolve())

                    out_file = output / f"{exp.fb_name}_接口变量表.xlsx"
                    wb = Workbook()
                    ws = wb.active
                    ws.title = exp.fb_name

                    cls._write_sheet(ws, exp)
                    wb.save(str(out_file))

                    exported_paths.append(str(out_file))
                    logger.info(f"已导出: {exp.fb_name} -> {out_file}")

            except Exception as e:
                logger.warning(f"跳过文件 {scl_file.name}: {e}")
                continue

        logger.info(f"批量导出完成: 共 {len(exported_paths)} 个FB/FC")
        return exported_paths

    @classmethod
    def export_check_report(cls, check_results: List, output_path: str) -> str:
        wb = Workbook()
        ws = wb.active
        ws.title = "检查报告"

        headers = ["序号", "规则ID", "严重级别", "文件路径", "行号", "描述", "修复建议"]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGNMENT
            cell.border = THIN_BORDER

        for row_idx, item in enumerate(check_results, start=2):
            if isinstance(item, dict):
                values = [
                    row_idx - 1,
                    item.get("rule_id", ""),
                    item.get("severity", ""),
                    item.get("file_path", ""),
                    item.get("line_number", ""),
                    item.get("message", ""),
                    item.get("suggestion", ""),
                ]
            else:
                values = [row_idx - 1] + [str(item)] * 6

            fill = EVEN_ROW_FILL if (row_idx % 2 == 0) else ODD_ROW_FILL
            for col_idx, val in enumerate(values, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.border = THIN_BORDER
                cell.fill = fill

        ws.freeze_panes = "A2"

        out = Path(output_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(out))

        logger.info(f"检查报告已导出: {out}")
        return str(out)
